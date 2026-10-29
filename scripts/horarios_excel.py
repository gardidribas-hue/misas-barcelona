"""Excel de revisión de horarios (data/horarios/Horarios_revision.xlsx).

Una hoja por distrito de Barcelona y por municipio (parroquias), y otra por tipo de templo. Columnas:
    A Nombre · B Web · C Horario de la fuente (automático) · D Horario revisado (Guillem) · E Fiabilidad (fórmula)
    F Fuente de la columna C · G Municipio · H ID
C es el horario que se ha leído (web propia; si no hay, Arzobispado u OpenStreetMap). D lo escribe Guillem a mano.
E compara C y D, y la web pública muestra el mismo aviso (horarios_app.py llama a estado()).
    python scripts/horarios_excel.py generar     reescribe el Excel con la columna C al día y conserva B y D
"""
import json, pathlib, re, sys, unicodedata

ROOT = pathlib.Path(__file__).resolve().parent.parent
H = ROOT / "data" / "horarios"
XLSX = H / "Horarios_revision.xlsx"
FUENTE = {"F1": "Web propia", "F3": "Arzobispado", "OSM": "OpenStreetMap", "F2": "misas.org"}
PARROQUIAS = ("Parroquia", "Parroquia personal", "Catedral")
HOJA_TIPO = {"Santuario": "Santuarios", "Basílica": "Basílicas", "Iglesia": "Iglesias", "Capilla": "Capillas",
             "Capilla de colegio": "Capillas", "Capilla de hospital": "Capillas de hospital",
             "Capilla de hospital (parroquia)": "Capillas de hospital", "Capilla de residencia": "Capillas de residencia",
             "Ermita": "Ermitas", "Monasterio / convento": "Monasterios y conventos", "Oratorio": "Oratorios"}
MUNICIPIO_HOJA = 4  # con menos parroquias, el municipio va a la hoja «Otros municipios»
CAB = ["Nombre", "Web (parroquia o comunidad pastoral)", "Horario leído (automático)", "Horario revisado (Guillem)",
       "Fiabilidad (automática)", "Fuente de la columna C", "Municipio", "ID"]

# Estados de la columna E: (código, texto que empieza la celda, color). La web usa el código.
ESTADOS = {
    "sin_horario": ("Sin horario · no sale en la web", "EDEDED"),
    "sin_web": ("Sin web · se muestra tu horario; puede cambiar, confírmalo", "FFF2CC"),
    "sin_lectura": ("Solo tu horario · su web no se ha podido leer; puede cambiar, confírmalo", "FFF2CC"),
    "sin_revisar": ("Sin revisar · fiabilidad media", "FFF2CC"),
    "sin_revisar_otra": ("Sin revisar · no es de la web de la parroquia, confírmalo", "FFF2CC"),
    "fiable": ("Fiable · el horario leído y tu revisión coinciden", "D9F2D9"),
    "no_coinciden": ("No coinciden · se muestra el de la web; revisa web y horario", "F8D7DA"),
    "revisado": ("Revisado · se muestra tu horario (la fuente dice otra cosa)", "FFF2CC"),
}


def igual(a, b):
    """La misma comparación que la fórmula de la columna E: sin espacios ni saltos de línea, sin mayúsculas, – = -."""
    f = lambda s: re.sub(r"[ \n\r]", "", (s or "")).replace("–", "-").lower()
    return f(a) == f(b)


def estado(c, d, fuente, web=""):
    """Código del estado y qué horario se muestra en la web: "C", "D" o None."""
    c, d = (c or "").strip(), (d or "").strip()
    if not c and not d:
        return "sin_horario", None
    if not c:
        return ("sin_lectura" if (web or "").strip() else "sin_web"), "D"
    if not d:
        return ("sin_revisar" if fuente == FUENTE["F1"] else "sin_revisar_otra"), "C"
    if igual(c, d):
        return "fiable", "C"
    # Manda la web de la parroquia; frente al Arzobispado u OpenStreetMap, manda la revisión.
    return ("no_coinciden", "C") if fuente == FUENTE["F1"] else ("revisado", "D")


def formula(i):
    n = lambda col: f'LOWER(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({col}{i}," ",""),CHAR(10),""),CHAR(13),""),"–","-"))'
    e = {k: f'"{v[0]}"' for k, v in ESTADOS.items()}
    web = f'F{i}="{FUENTE["F1"]}"'
    return (f'=IF(AND(TRIM(C{i})="",TRIM(D{i})=""),{e["sin_horario"]},'
            f'IF(TRIM(C{i})="",IF(TRIM(B{i})="",{e["sin_web"]},{e["sin_lectura"]}),'
            f'IF(TRIM(D{i})="",IF({web},{e["sin_revisar"]},{e["sin_revisar_otra"]}),'
            f'IF({n("C")}={n("D")},{e["fiable"]},'
            f'IF({web},{e["no_coinciden"]},{e["revisado"]})))))')


# --- De celebraciones a texto y de texto a celebraciones ------------------------------------------
ORD = ["lunes", "martes", "miercoles", "jueves", "viernes", "sabado", "domingo", "vispera_festivo", "festivo_precepto"]
AB = {"lunes": "Lun", "martes": "Mar", "miercoles": "Mié", "jueves": "Jue", "viernes": "Vie", "sabado": "Sáb",
      "domingo": "Dom", "vispera_festivo": "vísperas", "festivo_precepto": "festivos"}
TIPOS = [("misa", "Misas"), ("confesion", "Confesiones"), ("adoracion", "Adoración")]


def dias_txt(ds):
    ds = [d for d in ORD if d in ds]
    sem = [d for d in ds if d in ORD[:7]]
    extra = [AB[d] for d in ds if d not in ORD[:7]]
    if sem == ORD[:7]:
        base = "Cada día"
    elif len(sem) >= 3 and sem == ORD[ORD.index(sem[0]):ORD.index(sem[0]) + len(sem)]:
        base = f"{AB[sem[0]]}–{AB[sem[-1]]}"
    else:
        base = ", ".join(AB[d] for d in sem)
    if base and extra:
        return base + " y " + " y ".join(extra)
    return base or " y ".join(extra)


def hora(h):
    return re.sub(r"^0(\d)", r"\1", h)


def horario_txt(cels, md):
    """«Misas: Lun–Vie 8:00, 19:00; Sáb y vísperas 20:00; Dom y festivos 11:00» (una línea por tipo, temporada de hoy)."""
    from horarios_app import en_temporada
    vig = [c for c in cels if en_temporada(c.get("temporada"), md)
           and not ((c.get("temporada") or {}).get("nombre") in ("Semana Santa", "Navidad") and not (c.get("temporada") or {}).get("inicio"))]
    lineas = []
    for tipo, nom in TIPOS:
        grupos, libres = {}, []
        for c in (c for c in vig if c["tipo"] == tipo):
            if c.get("relativa_a_misa"):
                libres.append("antes de misa")
            elif c.get("a_peticion"):
                libres.append("a petición")
            elif c.get("inicio") and c.get("dias") and not c.get("semanas_mes"):
                grupos.setdefault(tuple(d for d in ORD if d in c["dias"]), set()).add(hora(c["inicio"]) + (f"–{hora(c['fin'])}" if c.get("fin") else ""))
        clave_h = lambda h: tuple(int(x) for x in re.findall(r"\d+", h)[:2])
        partes = [f"{dias_txt(k)} {', '.join(sorted(v, key=clave_h))}" for k, v in sorted(grupos.items(), key=lambda kv: [ORD.index(d) for d in kv[0]])]
        partes += sorted(set(libres))
        if partes:
            lineas.append(f"{nom}: " + "; ".join(partes))
    return "\n".join(lineas)


def _sin_acentos(s):
    return unicodedata.normalize("NFKD", s.replace("–", "-")).encode("ascii", "ignore").decode().lower()


DIA_PAL = {"lun": "lunes", "dl": "lunes", "dill": "lunes", "mar": "martes", "dt": "martes", "dim": "martes",
           "mie": "miercoles", "mier": "miercoles", "dc": "miercoles", "dime": "miercoles", "jue": "jueves", "dj": "jueves",
           "dij": "jueves", "vie": "viernes", "dv": "viernes", "div": "viernes", "sab": "sabado", "ds": "sabado",
           "diss": "sabado", "dom": "domingo", "dg": "domingo", "dium": "domingo"}
TIPO_PAL = {"misa": "misa", "misas": "misa", "missa": "misa", "misses": "misa", "confesion": "confesion",
            "confesiones": "confesion", "confessio": "confesion", "confessions": "confesion", "adoracion": "adoracion",
            "adoracio": "adoracion", "exposicion": "adoracion"}
HORA_RE = re.compile(r"(\d{1,2})(?:[:.h](\d{2}))?\s*h?(?:\s*[-–a]\s*(\d{1,2})(?:[:.h](\d{2}))?\s*h?)?")


def _dia(pal):
    p = _sin_acentos(pal).strip(".")
    if p.startswith(("festiv", "fest")):
        return "festivo_precepto"
    if p.startswith(("vispera", "vigil", "vespr")):
        return "vispera_festivo"
    for k in sorted(DIA_PAL, key=len, reverse=True):
        if p.startswith(k):
            return DIA_PAL[k]
    return None


def dias_de(txt):
    t = _sin_acentos(txt)
    if re.search(r"cada dia|todos los dias|diari", t):
        return ORD[:7] + [d for d in ("festivo_precepto",) if "festiv" in t]
    if re.search(r"laborables|feiners", t):
        base = ORD[:5]
    else:
        base = []
    out = list(base)
    m = re.search(r"([a-z.]+)(?:\s*-\s*|\s+a\s+)([a-z.]+)", t)
    if m and _dia(m.group(1)) in ORD[:7] and _dia(m.group(2)) in ORD[:7]:
        a, b = ORD.index(_dia(m.group(1))), ORD.index(_dia(m.group(2)))
        out += ORD[a:b + 1]
        t = t[:m.start()] + " " + t[m.end():]
    for pal in re.findall(r"[a-z.]+", t):
        d = _dia(pal)
        if d and d not in out and pal not in ("a", "y", "i", "de"):
            out.append(d)
    return [d for d in ORD if d in out]


def leer_horario(texto):
    """Texto de la columna D → (celebraciones, textos libres por tipo). Lo que no se entiende va como texto."""
    cels, libres = [], {}
    for linea in re.split(r"[\n\r]+", texto or ""):
        if not linea.strip():
            continue
        tipo = "misa"
        m = re.match(r"\s*([^\W\d_]+)\s*:\s*(.*)$", linea)
        if m and _sin_acentos(m.group(1)) in TIPO_PAL:
            tipo, linea = TIPO_PAL[_sin_acentos(m.group(1))], m.group(2)
        for grupo in re.split(r";", linea):
            grupo = grupo.strip()
            if not grupo:
                continue
            m = re.search(r"\d", grupo)
            dias = dias_de(grupo[:m.start()]) if m else []
            horas = [h for h in HORA_RE.finditer(grupo[m.start():])] if m else []
            if not dias or not horas:
                libres.setdefault(tipo, []).append(grupo[0].upper() + grupo[1:])
                continue
            for h in horas:
                ini = f"{int(h.group(1)):02d}:{h.group(2) or '00'}"
                fin = f"{int(h.group(3)):02d}:{h.group(4) or '00'}" if h.group(3) else None
                cels.append({"tipo": tipo, "dias": dias, "inicio": ini, "fin": fin})
    return cels, libres


# --- Hojas -----------------------------------------------------------------------------------------
def hoja_de(lug, distrito, n_muni):
    if lug["tipo"] in PARROQUIAS:
        if lug["municipio"] == "Barcelona":
            return f"BCN {distrito}" if distrito else "BCN sin distrito"
        return lug["municipio"] if n_muni.get(lug["municipio"], 0) >= MUNICIPIO_HOJA else "Otros municipios"
    return HOJA_TIPO.get(lug["tipo"], "Otros templos")


def web_por_defecto(lug):
    return ((lug.get("fuentes") or {}).get("F1") or {}).get("url") or lug.get("web") or ""


def leer(xlsx=XLSX):
    """{ID: {"web": B, "revisado": D}} de todas las hojas del Excel de revisión (vacío si no existe)."""
    if not pathlib.Path(xlsx).exists():
        return {}
    import openpyxl
    out = {}
    for ws in openpyxl.load_workbook(xlsx, data_only=True).worksheets:
        filas = ws.iter_rows(values_only=True)
        cab = [str(x or "").strip() for x in next(filas, [])]
        if "ID" not in cab:
            continue
        i_id, i_web, i_rev = cab.index("ID"), cab.index(CAB[1]) if CAB[1] in cab else 1, cab.index(CAB[3]) if CAB[3] in cab else 3
        for f in filas:
            tid = f[i_id] if len(f) > i_id else None
            if tid:
                out[str(tid).strip()] = {"web": str(f[i_web] or "").strip(), "revisado": str(f[i_rev] or "").strip()}
    return out


def webs_cambiadas(lugares, rev=None):
    """Webs que Guillem ha cambiado en la columna B: {ID: URL}. horarios.py las usa como F1."""
    rev = leer() if rev is None else rev
    return {l["id"]: rev[l["id"]]["web"] for l in lugares
            if l["id"] in rev and rev[l["id"]]["web"].startswith("http") and rev[l["id"]]["web"] != web_por_defecto(l)}


def fila(lug, t, rev, md):
    c = horario_txt(t.get("celebraciones") or [], md) if t.get("estado") in ("ok", "pendiente de verificar") else ""
    fuente = FUENTE.get(t.get("fuente"), "") if c else ""
    r = rev.get(lug["id"], {})
    return {"nombre": lug["nombre"], "web": r.get("web") or web_por_defecto(lug), "c": c, "d": r.get("revisado", ""),
            "fuente": fuente, "municipio": lug["municipio"], "id": lug["id"]}


def generar(xlsx=XLSX, fecha=None):
    import datetime as dt
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.formatting.rule import FormulaRule
    from horarios_app import distrito_de
    md = (fecha or dt.date.today().isoformat())[5:]
    lugares = json.loads((H / "lugares.json").read_text("utf-8"))["lugares"]
    templos = json.loads((H / "horarios.json").read_text("utf-8"))["templos"]
    geo_js = (ROOT / "data" / "geo.js").read_text("utf-8")
    geo = json.loads(geo_js[geo_js.index("{"):geo_js.rindex("}") + 1])
    rev = leer(xlsx)
    n_muni = {}
    for l in lugares:
        if l["tipo"] in PARROQUIAS:
            n_muni[l["municipio"]] = n_muni.get(l["municipio"], 0) + 1
    hojas = {}
    for l in lugares:
        dist = distrito_de(l["lat"], l["lng"], geo) if l["tipo"] in PARROQUIAS and l["municipio"] == "Barcelona" and l.get("lat") else None
        hojas.setdefault(hoja_de(l, dist, n_muni), []).append(fila(l, templos.get(l["id"], {}), rev, md))
    tipos = list(dict.fromkeys(HOJA_TIPO.values())) + ["Otros templos"]
    orden = (sorted(h for h in hojas if h.startswith("BCN ")) + sorted((h for h in hojas if not h.startswith("BCN ") and h not in tipos and h != "Otros municipios"), key=_sin_acentos)
             + (["Otros municipios"] if "Otros municipios" in hojas else []) + [h for h in tipos if h in hojas])

    F = "Arial"
    thin = Side(style="thin", color="D9D9D9")
    wb = Workbook()
    ex = wb.active
    ex.title = "Cómo funciona"
    txt = [
        ("Horarios de misas, confesiones y adoración · revisión", True, None), ("", False, None),
        ("Una hoja por distrito de Barcelona y por municipio (parroquias); los municipios con menos de "
         f"{MUNICIPIO_HOJA} parroquias van juntos en «Otros municipios». Después, una hoja por tipo de templo. Dentro de cada hoja, por orden alfabético.", False, None),
        ("", False, None), ("Columnas", True, None),
        ("A · Nombre del templo.", False, None),
        ("B · Web de la parroquia o de su comunidad pastoral. Si la cambias, el programa leerá el horario de la web nueva el lunes siguiente.", False, None),
        ("C · Horario leído: lo rellena el programa cada lunes. Si la parroquia no tiene web propia (o no se puede leer), es el del Arzobispado. No lo toques.", False, None),
        ("D · Horario revisado: el tuyo. Escríbelo con el mismo formato que la columna C (lo más fácil es copiar C y corregir).", False, None),
        ("E · Fiabilidad: automática, no la toques. Cambia al momento cuando editas D, y la web pública muestra el mismo aviso.", False, None),
        ("F · De dónde sale la columna C: Web propia, Arzobispado u OpenStreetMap. G · Municipio. H · ID: código fijo del templo, no lo cambies.", False, None),
        ("", False, None), ("Fiabilidad (columna E) y qué sale en la web pública", True, None),
        ("Verde · Fiable: el horario leído y tu revisión coinciden.", False, ESTADOS["fiable"][1]),
        ("Rojo · No coinciden: la web de la parroquia dice otra cosa que tu revisión. Se muestra el de la web, con aviso de que no es fiable.", False, ESTADOS["no_coinciden"][1]),
        ("Amarillo · Sin revisar: hay horario leído pero aún no lo has revisado. Se muestra con fiabilidad media (o, si viene del Arzobispado, pidiendo confirmarlo).", False, ESTADOS["sin_revisar"][1]),
        ("Amarillo · Sin web: solo está tu horario. Se muestra el tuyo, avisando de que la parroquia no tiene web y puede cambiar.", False, ESTADOS["sin_web"][1]),
        ("Amarillo · Solo tu horario: tiene web pero no se ha podido leer. Se muestra el tuyo, avisando de que puede cambiar.", False, ESTADOS["sin_lectura"][1]),
        ("Amarillo · Revisado: el Arzobispado dice otra cosa que tu revisión. Se muestra el tuyo, porque frente al Arzobispado manda tu revisión.", False, ESTADOS["revisado"][1]),
        ("Gris · Sin horario: ni horario leído ni revisión. No sale en la web pública.", False, ESTADOS["sin_horario"][1]),
        ("", False, None), ("Formato del horario", True, None),
        ("Una línea por tipo: «Misas: …», «Confesiones: …», «Adoración: …». Grupos separados por punto y coma.", False, None),
        ("Días: Lun, Mar, Mié, Jue, Vie, Sáb, Dom; intervalos con guion (Lun–Vie); «Dom y festivos», «Sáb y vísperas», «Cada día». Horas en 24 h separadas por comas; un intervalo con guion (18:00–19:00).", False, None),
        ("Ejemplo: Misas: Lun–Vie 8:00, 19:00; Sáb y vísperas 20:00; Dom y festivos 11:00, 12:30", False, None),
        ("", False, None), ("Cómo se actualiza la web", True, None),
        ("Sube el Excel cambiado a GitHub (data/horarios/Horarios_revision.xlsx, con el mismo nombre) y la web se actualiza en unos minutos. "
         "Cada lunes el programa vuelve a leer las webs, reescribe la columna C y conserva lo que hayas puesto en B y D.", False, None),
    ]
    for i, (t_, b, color) in enumerate(txt, 1):
        c = ex.cell(row=i, column=1, value=t_)
        c.font = Font(name=F, bold=b, size=14 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if color:
            c.fill = PatternFill("solid", fgColor=color)
    ex.column_dimensions["A"].width = 130

    resumen = []
    for nombre in orden:
        filas = sorted(hojas[nombre], key=lambda r: (_sin_acentos(r["municipio"]) if nombre == "Otros municipios" else "", _sin_acentos(r["nombre"])))
        ws = wb.create_sheet(nombre[:31])
        for j, h in enumerate(CAB, 1):
            c = ws.cell(row=1, column=j, value=h)
            c.font = Font(name=F, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="7F3B3B")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for i, r in enumerate(filas, 2):
            vals = [r["nombre"], r["web"] or None, r["c"] or None, r["d"] or None, formula(i), r["fuente"] or None, r["municipio"], r["id"]]
            for j, v in enumerate(vals, 1):
                c = ws.cell(row=i, column=j, value=v)
                c.font = Font(name=F, color="999999") if j in (6, 7, 8) else Font(name=F)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.border = Border(bottom=thin)
            if r["web"].startswith("http"):
                ws.cell(row=i, column=2).hyperlink = r["web"]
                ws.cell(row=i, column=2).font = Font(name=F, color="0563C1", underline="single")
            ws.cell(row=i, column=3).fill = PatternFill("solid", fgColor="F5F5F5")
        last = len(filas) + 1
        for codigo, (texto, color) in ESTADOS.items():
            pref = texto.split(" · ")[0]
            ws.conditional_formatting.add(f"E2:E{last}", FormulaRule(formula=[f'LEFT(E2,{len(pref)})="{pref}"'], fill=PatternFill("solid", fgColor=color)))
        for col, w in zip("ABCDEFGH", (34, 34, 46, 46, 30, 14, 16, 9)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:H{last}"
        resumen.append((nombre, len(filas)))
    wb.active = 1
    wb.save(xlsx)
    print(f"{sum(n for _, n in resumen)} templos en {len(resumen)} hojas: " + ", ".join(f"{h} {n}" for h, n in resumen))


if __name__ == "__main__":
    sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
    if sys.argv[1:2] == ["generar"]:
        generar(sys.argv[2] if len(sys.argv) > 2 else XLSX)
    else:
        sys.exit(__doc__)
